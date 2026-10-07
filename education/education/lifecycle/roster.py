"""Read the class lists staff prepare for a new year.

A roster maps School ID -> Section. Two layouts are understood, in .xlsx or
.csv:

* one sheet with a School ID column and a Section column (the audit
  workbook's "All students" sheet is exactly this); and
* the class-list workbooks themselves, one sheet per section, where the
  sheet's name is the section ("LKG A", "Grade 10A").

Every sheet of a workbook is read, so a file can mix both.
"""

import io

import frappe
from frappe import _

from education.education.lifecycle.common import clean_spaces, normalize_school_id

SCHOOL_ID_HEADERS = {
	"schoolid",
	"rosterschoolid",
	"studentid",
	"schoolidnumber",
	"id",
	"idno",
	"idnumber",
}
SECTION_HEADERS = {"section", "studentgroup", "class", "group", "homeroom", "classsection"}
NAME_HEADERS = {"studentname", "name", "fullname", "rostername"}

# How far down a sheet to look for its header row.
HEADER_SCAN_ROWS = 10


def _header_key(value):
	return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def _find_columns(rows):
	"""Return (header_row_index, id_col, section_col, name_col) or None."""
	for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
		keys = [_header_key(cell) for cell in row]
		id_col = next((i for i, k in enumerate(keys) if k in SCHOOL_ID_HEADERS), None)
		if id_col is None:
			continue
		section_col = next((i for i, k in enumerate(keys) if k in SECTION_HEADERS), None)
		name_col = next((i for i, k in enumerate(keys) if k in NAME_HEADERS), None)
		return index, id_col, section_col, name_col
	return None


def _sheets_from_file(file_url):
	"""Yield (sheet_name, rows) for every sheet in an attached roster file."""
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	content = file_doc.get_content()
	name = (file_doc.file_name or file_url).lower()

	if name.endswith(".csv"):
		from frappe.utils.csvutils import read_csv_content

		if isinstance(content, bytes):
			content = content.decode("utf-8-sig", errors="replace")
		yield "", read_csv_content(content)
		return

	if not name.endswith((".xlsx", ".xlsm")):
		frappe.throw(_("Upload the roster as an .xlsx or .csv file."))

	from openpyxl import load_workbook

	if isinstance(content, str):
		content = content.encode()
	workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	try:
		for sheet in workbook.worksheets:
			yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
	finally:
		workbook.close()


def read_roster(file_url):
	"""Parse a roster file.

	Returns a dict::

	    {
	        "sections": {normalized_school_id: section},
	        "names": {normalized_school_id: name as written on the roster},
	        "conflicts": {normalized_school_id: [section, section, ...]},
	        "rows": int,          # rows that carried a School ID
	        "sheets": [sheet names that were used],
	    }

	A School ID listed under two different sections is a conflict: it is left
	out of ``sections`` so nobody is silently placed in the wrong class.
	"""
	seen = {}
	names = {}
	rows_read = 0
	sheets_used = []

	for sheet_name, rows in _sheets_from_file(file_url):
		found = _find_columns(rows)
		if not found:
			continue
		header_index, id_col, section_col, name_col = found
		if section_col is None and not sheet_name:
			continue

		used = False
		for row in rows[header_index + 1 :]:
			if id_col >= len(row):
				continue
			school_id = normalize_school_id(row[id_col])
			if not school_id:
				continue
			section = row[section_col] if section_col is not None and section_col < len(row) else sheet_name
			section = clean_spaces(section)
			if not section:
				continue
			rows_read += 1
			used = True
			seen.setdefault(school_id, set()).add(section)
			if name_col is not None and name_col < len(row) and row[name_col]:
				names.setdefault(school_id, clean_spaces(row[name_col]))
		if used:
			sheets_used.append(sheet_name or _("(single sheet)"))

	sections, conflicts = {}, {}
	for school_id, values in seen.items():
		# "Grade 10A" and "Grade 10 A" are the same class written two ways.
		distinct = {v.replace(" ", "").lower(): v for v in values}
		if len(distinct) == 1:
			sections[school_id] = next(iter(values))
		else:
			conflicts[school_id] = sorted(values)

	return frappe._dict(
		sections=sections,
		names=names,
		conflicts=conflicts,
		rows=rows_read,
		sheets=sheets_used,
	)
